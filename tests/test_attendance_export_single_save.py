"""Eksport zachowuje dane i ustawienia druku przy jednym zapisie XLSX."""
from unittest.mock import Mock

import openpyxl
import pytest

import profile_attendance_export_runtime as export
import profile_employee_editor_finish_runtime as finish


@pytest.mark.parametrize("has_rows", [False, True])
def test_export_saves_once_with_print_settings(monkeypatch, tmp_path, has_rows):
    rows = [{
        "date": "2026-09-07", "slot": "RANO", "first_login": "05:57:00",
        "status": "Obecny", "day_value": 1.0, "absence": "ŚW",
        "overtime_hours": 2.0, "overtime_type": "zwykle",
        "source": "foreman", "note": "Próba eksportu",
    }] if has_rows else []
    monkeypatch.setattr(export, "_attendance_export_rows", lambda *_: rows)
    monkeypatch.setattr(finish, "_attendance_export_rows", lambda *_: rows)
    monkeypatch.setattr(finish, "_wm_attendance_export_opt_v1", False, raising=False)
    original_export = finish._export_attendance_xlsx
    export.install()
    assert finish._export_attendance_xlsx is original_export
    monkeypatch.setattr(finish.workforce_profile_service, "get_user", lambda _: {})
    monkeypatch.setattr(finish.workforce_profile_service, "display_name", lambda _: "Jan Żółć")
    monkeypatch.setattr(finish.attendance_service, "summary_for_month", lambda *_: {
        "days": int(has_rows), "overtime_hours": 2 * int(has_rows),
    })
    saved = []
    original_save = openpyxl.Workbook.save
    original_load = openpyxl.load_workbook

    def save(wb, path):
        saved.append(path)
        assert wb["Ewidencja"].page_setup.orientation == "landscape"
        original_save(wb, path)

    read = Mock(side_effect=AssertionError("Eksport nie powinien odczytywać XLSX"))
    monkeypatch.setattr(openpyxl.Workbook, "save", save)
    monkeypatch.setattr(openpyxl, "load_workbook", read)
    path = tmp_path / "obecnosc.xlsx"
    assert finish._export_attendance_xlsx(path, "jan", 2026, 9) == path
    assert saved == [path]
    read.assert_not_called()

    wb = original_load(path)
    try:
        assert wb.sheetnames == ["Ewidencja", "Podsumowanie"]
        ws, sm = wb.worksheets
        assert ws["B1"].value == "Jan Żółć"
        assert ws["B2"].value == "jan"
        assert ws["B3"].value == "2026-09"
        assert list(ws.values)[4] == (
            "Data", "Zmiana", "Pierwsze logowanie", "Status", "Dniówka",
            "Nieobecność", "Nadgodziny [h]", "Typ nadgodzin", "Źródło", "Uwagi",
        )
        assert list(ws.values)[5:] == ([tuple(rows[0].values())] if has_rows else [])
        assert list(sm.values)[4:] == [
            ("Dniówki", int(has_rows)), ("Soboty", 0),
            ("Nadgodziny [h]", 2 * int(has_rows)), ("L4 [dni]", 0),
            ("ŚW [dni]", int(has_rows)), ("NN [dni]", 0),
            ("UR [dni]", 0), ("UŻ [dni]", 0), ("Braki", 0), ("Do decyzji", 0),
        ]
        assert ws.freeze_panes == "A6"
        assert ws.auto_filter.ref == f"A5:J{ws.max_row}"
        assert ws.print_title_rows == "$5:$5"
        assert ws.print_area == f"'Ewidencja'!$A$1:$J${ws.max_row}"
        assert sm.print_area == f"'Podsumowanie'!$A$1:$B${sm.max_row}"
        for sheet, orientation in [(ws, "landscape"), (sm, "portrait")]:
            assert sheet.page_setup.orientation == orientation
            assert sheet.page_setup.fitToWidth == 1
            assert sheet.page_setup.fitToHeight == 0
            assert sheet.sheet_properties.pageSetUpPr.fitToPage
            assert sheet.print_options.horizontalCentered
        assert ws.page_margins.left == ws.page_margins.right == 0.25
        assert ws.page_margins.top == ws.page_margins.bottom == 0.5
    finally:
        wb.close()
