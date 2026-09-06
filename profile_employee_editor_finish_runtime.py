# version: 1.1
"""Końcowe dopracowanie edytora pracownika bez zmiany źródeł danych."""
from __future__ import annotations

import tkinter as tk
from calendar import monthrange
from datetime import date
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Any

from services import attendance_service, leave_balance_service, workforce_profile_service
from ui_context_help import add_help_button

_INSTALLED = False
_OPEN_EMPLOYEE_WINDOWS: dict[tuple[int, str], tk.Toplevel] = {}


def _walk(widget):
    out = []
    try:
        children = widget.winfo_children()
    except Exception:
        children = []
    for child in children:
        out.append(child)
        out.extend(_walk(child))
    return out


def _all_toplevels(root) -> set[tk.Toplevel]:
    return {widget for widget in _walk(root) if isinstance(widget, tk.Toplevel)}


def _employee_key(owner, login: str) -> tuple[int, str]:
    try:
        root = owner.winfo_toplevel()
    except Exception:
        root = owner
    user = workforce_profile_service.get_user(login) or {}
    stable = str(user.get("user_id") or login or "").strip().casefold()
    return id(root), stable


def _resolve_login(user_id: str, fallback: str) -> str:
    uid = str(user_id or "").strip().casefold()
    if uid:
        try:
            for row in workforce_profile_service.list_users():
                if str(row.get("user_id") or "").strip().casefold() == uid:
                    current = str(row.get("login") or "").strip()
                    if current:
                        return current
        except Exception:
            pass
    return str(fallback or "").strip()


def _notebook(win) -> ttk.Notebook | None:
    for widget in _walk(win):
        if isinstance(widget, ttk.Notebook):
            return widget
    return None


def _tabs(nb: ttk.Notebook) -> dict[str, str]:
    out: dict[str, str] = {}
    try:
        for tab_id in nb.tabs():
            out[str(nb.tab(tab_id, "text"))] = tab_id
    except Exception:
        pass
    return out


def _select_tab(win, name: str) -> None:
    nb = _notebook(win)
    if nb is None:
        return
    requested = str(name or "")
    requested = "Podstawowe" if requested == "Dane" else requested
    tab_id = _tabs(nb).get(requested)
    if tab_id:
        try:
            if requested in {"Historia", "Uprawnienia"}:
                nb.add(tab_id, text=requested)
            nb.select(tab_id)
        except Exception:
            pass


def _raise_window(win, initial_tab: str) -> bool:
    try:
        if not win.winfo_exists():
            return False
        win.deiconify()
        win.lift()
        _select_tab(win, initial_tab)
        try:
            win.focus_force()
        except Exception:
            pass
        return True
    except Exception:
        return False


def _fmt(value: Any) -> str:
    if value is None or value == "":
        return "—"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _compact(value: Any) -> str:
    if isinstance(value, dict):
        hours = value.get("hours")
        kind = value.get("type")
        status = value.get("status")
        parts = []
        if hours not in (None, "", 0, 0.0):
            parts.append(f"{_fmt(hours)} h")
        if kind:
            parts.append(str(kind))
        if status:
            parts.append(str(status))
        return ", ".join(parts) or "—"
    return _fmt(value)


def _changes(before: Any, after: Any) -> list[str]:
    if not isinstance(before, dict) or not isinstance(after, dict):
        return []
    labels = {
        "slot": "zmiana",
        "day_value": "dniówka",
        "reason": "nieobecność",
        "absence_code": "nieobecność",
        "pay_code": "kod dnia",
        "overtime": "nadgodziny",
        "entitlement": "należny",
        "adjustment": "korekta",
        "carryover": "zaległy",
        "remaining": "pozostało",
        "login": "login",
        "imie": "imię",
        "nazwisko": "nazwisko",
        "rola": "rola",
        "status": "status",
        "zatrudniony_od": "zatrudniony od",
        "zatrudniony_do": "zatrudniony do",
        "telefon": "telefon",
        "email": "e-mail",
        "tryb_zmian": "grafik",
        "shift_mode": "grafik",
        "shift_anchor": "kotwica",
    }
    out: list[str] = []
    seen_labels: set[str] = set()
    for key, label in labels.items():
        if key not in before and key not in after:
            continue
        left, right = before.get(key), after.get(key)
        if left == right or label in seen_labels:
            continue
        seen_labels.add(label)
        out.append(f"{label}: {_compact(left)} → {_compact(right)}")
    return out


def _human_action(row: dict) -> tuple[str, str]:
    action = str(row.get("action") or row.get("_source") or "").strip()
    labels = {
        "profil": "Dane profilu",
        "reset_pin": "Reset PIN",
        "obecnosc": "Korekta obecności",
        "manual_day": "Korekta dniówki",
        "manual_day_replace_absence": "Nieobecność → dniówka",
        "manual_day_move": "Zmiana zmiany",
        "absence_kind_edit": "Zmiana nieobecności",
        "overtime": "Nadgodziny",
        "saldo_urlopu": "Saldo urlopu",
        "grafik": "Grafik",
        "uprawnienia": "Uprawnienia",
    }
    title = labels.get(action, action.replace("_", " ").strip().capitalize() or "Zmiana")
    note = str(row.get("note") or "").strip()
    changes = _changes(row.get("before"), row.get("after"))
    detail = "; ".join(changes)
    if note and note not in detail:
        detail = f"{detail}; {note}".strip("; ")
    if not detail:
        day = str(row.get("date") or "").strip()
        slot = str(row.get("slot") or "").strip()
        detail = " • ".join(value for value in (day, slot) if value) or "Zmiana zapisana w WM"
    return title, detail


def _history_rows(login: str) -> list[dict]:
    try:
        import profile_foreman_edit_runtime as edit_runtime
        profile_rows = edit_runtime._profile_audit(login, 250)
    except Exception:
        profile_rows = []
    rows: list[dict] = []
    for row in profile_rows:
        item = dict(row)
        item["_source"] = "profil"
        rows.append(item)
    try:
        for row in attendance_service.audit_for_login(login, 250):
            item = dict(row)
            item["_source"] = "obecność"
            rows.append(item)
    except Exception:
        pass
    rows.sort(key=lambda item: str(item.get("ts") or ""), reverse=True)
    return rows[:400]


def _decorate_history(win, nb: ttk.Notebook, login: str) -> None:
    history_id = _tabs(nb).get("Historia")
    if not history_id:
        return
    try:
        history = win.nametowidget(history_id)
    except Exception:
        return
    tree = next((widget for widget in _walk(history) if isinstance(widget, ttk.Treeview)), None)
    if tree is None:
        return
    try:
        tree.heading("action", text="Co zmieniono")
        tree.heading("actor", text="Kto")
        tree.heading("note", text="Szczegóły")
        tree.column("action", width=175, anchor="w")
        tree.column("note", width=500, anchor="w")
    except Exception:
        pass

    user_id = str((workforce_profile_service.get_user(login) or {}).get("user_id") or "")

    def refresh() -> None:
        try:
            if not win.winfo_exists():
                return
            tree.delete(*tree.get_children())
            for row in _history_rows(_resolve_login(user_id, login)):
                title, detail = _human_action(row)
                ts = str(row.get("ts") or "").replace("T", " ")[:19]
                tree.insert("", "end", values=(ts, title, row.get("actor") or "—", detail))
        except Exception:
            pass

    for widget in _walk(history):
        if isinstance(widget, ttk.Button):
            try:
                if str(widget.cget("text")) == "Odśwież historię":
                    widget.configure(command=refresh)
            except Exception:
                pass
    refresh()

    try:
        root = win.winfo_toplevel()
        bind_id = root.bind("<<ProfileDataUpdated>>", lambda _e: win.after_idle(refresh), add="+")
    except Exception:
        root, bind_id = None, None

    if root is not None and bind_id:
        def cleanup(event) -> None:
            if event.widget is win:
                try:
                    root.unbind("<<ProfileDataUpdated>>", bind_id)
                except Exception:
                    pass
        win.bind("<Destroy>", cleanup, add="+")


def _grid_widget(parent, row: int, cls) -> Any | None:
    for widget in parent.winfo_children():
        if not isinstance(widget, cls):
            continue
        try:
            if int(widget.grid_info().get("row", -1)) == row:
                return widget
        except Exception:
            pass
    return None


def _set_textvariable(parent, widget, value: Any) -> None:
    try:
        name = str(widget.cget("textvariable") or "")
        if name:
            parent.setvar(name, str(value))
    except Exception:
        pass


def _decorate_leave_year(win, nb: ttk.Notebook, login: str) -> None:
    leaves_id = _tabs(nb).get("Urlopy")
    if not leaves_id:
        return
    try:
        leaves = win.nametowidget(leaves_id)
    except Exception:
        return
    year_entry = _grid_widget(leaves, 0, ttk.Entry)
    ent_entry = _grid_widget(leaves, 1, ttk.Entry)
    adj_entry = _grid_widget(leaves, 2, ttk.Entry)
    carry_entry = _grid_widget(leaves, 3, ttk.Entry)
    status_label = _grid_widget(leaves, 4, ttk.Label)
    if year_entry is None or ent_entry is None or adj_entry is None or carry_entry is None:
        return

    user_id = str((workforce_profile_service.get_user(login) or {}).get("user_id") or "")
    try:
        var_name = str(year_entry.cget("textvariable") or "")
        current = date.today().year
        years = tuple(str(year) for year in range(current - 1, current + 2))
        year_entry.grid_remove()
        combo = ttk.Combobox(leaves, textvariable=var_name, values=years, state="readonly", width=10)
        combo.grid(row=0, column=1, sticky="ew", pady=4)
        add_help_button(
            leaves,
            "Wybierz rok rozliczenia urlopu. WM pokaże należny, zaległy, wykorzystany, oczekujący i pozostały urlop dla wybranego roku.",
            row=0, column=2, padx=(6, 0), sticky="w",
        )
    except Exception:
        return

    def load_year(_event=None) -> None:
        try:
            year = int(leaves.getvar(var_name))
            current_login = _resolve_login(user_id, login)
            balance = leave_balance_service.get_balance(current_login, year)
            import profile_foreman_edit_runtime as edit_runtime
            _set_textvariable(leaves, ent_entry, f"{float(balance.get('entitlement') or 0):g}")
            _set_textvariable(leaves, adj_entry, f"{float(balance.get('adjustment') or 0):g}")
            _set_textvariable(leaves, carry_entry, edit_runtime._format_carryover(balance, year))
            if status_label is not None:
                text = (
                    f"🟡 Zaległy: {float(balance.get('carryover') or 0):g}   "
                    f"🔵 Wykorzystano: {float(balance.get('used') or 0):g}   "
                    f"🟠 Oczekuje: {float(balance.get('pending') or 0):g}   "
                    f"🟢 Pozostało: {float(balance.get('remaining') or 0):g}"
                )
                _set_textvariable(leaves, status_label, text)
        except Exception:
            pass

    combo.bind("<<ComboboxSelected>>", load_year, add="+")
    load_year()


def _month_choices() -> tuple[str, ...]:
    today = date.today()
    current_index = today.year * 12 + today.month - 1
    values = []
    for offset in range(2, -13, -1):
        index = current_index + offset
        year, month0 = divmod(index, 12)
        values.append(f"{year:04d}-{month0 + 1:02d}")
    return tuple(values)


def _month_tuple(value: str) -> tuple[int, int]:
    text = str(value or "").strip()
    year_s, month_s = text.split("-", 1)
    year, month = int(year_s), int(month_s)
    if month < 1 or month > 12:
        raise ValueError("Nieprawidłowy miesiąc.")
    return year, month


def _absence_code(value: Any) -> str:
    raw = str(value or "").strip().upper().replace("-", "_").replace(" ", "_")
    mapping = {
        "SW": "ŚW", "ŚW": "ŚW", "SILA_WYZSZA": "ŚW", "SIŁA_WYŻSZA": "ŚW",
        "SILA_WYZSZA_50": "ŚW", "URLOP": "UR", "UR": "UR",
        "URLOP_WYPOCZYNKOWY": "UR", "UZ": "UŻ", "UŻ": "UŻ",
        "URLOP_NA_ZADANIE": "UŻ", "L4": "L4", "NN": "NN",
    }
    return mapping.get(raw, raw)


def _absence_codes(login: str, day_text: str, row: dict | None = None) -> list[str]:
    out: list[str] = []
    reason = _absence_code((row or {}).get("reason"))
    if reason:
        out.append(reason)
    try:
        from services import leave_workflow_service
        leaves = leave_workflow_service.active_absences_for_day(login, day_text)
    except Exception:
        leaves = []
    for leave in leaves:
        code = _absence_code(leave.get("type"))
        if code and code not in out:
            out.append(code)
    return out


def _attendance_status_text(row: dict) -> str:
    status = str(row.get("status") or "")
    labels = {
        attendance_service.STATUS_PRESENT: "Obecny",
        attendance_service.STATUS_PENDING_LATE: "Późne logowanie",
        attendance_service.STATUS_MISSING: "Brak",
        attendance_service.STATUS_EXCUSED: "Nieobecność",
        attendance_service.STATUS_SATURDAY: "Sobota",
        attendance_service.STATUS_PLANNED: "Plan",
    }
    return labels.get(status, status or "—")


def _attendance_export_rows(login: str, year: int, month: int) -> list[dict]:
    records = [dict(row) for row in attendance_service.month_records(login, year, month)]
    existing_days = {str(row.get("date") or "")[:10] for row in records}
    for day_no in range(1, monthrange(year, month)[1] + 1):
        day_text = date(year, month, day_no).isoformat()
        if day_text not in existing_days and _absence_codes(login, day_text):
            records.append({
                "date": day_text, "slot": "", "status": attendance_service.STATUS_EXCUSED,
                "day_value": 0.0, "reason": "", "synthetic": True,
            })

    out: list[dict] = []
    for row in records:
        day_text = str(row.get("date") or "")[:10]
        overtime = row.get("overtime") if isinstance(row.get("overtime"), dict) else {}
        first_login = str(row.get("first_login_ts") or row.get("logged_ts") or "")
        if "T" in first_login:
            first_login = first_login.split("T", 1)[1][:8]
        codes = _absence_codes(login, day_text, row)
        out.append({
            "date": day_text,
            "slot": str(row.get("slot") or "—"),
            "first_login": first_login or "—",
            "status": _attendance_status_text(row),
            "day_value": float(row.get("day_value") or 0.0),
            "absence": ", ".join(codes),
            "overtime_hours": float(overtime.get("hours") or 0.0),
            "overtime_type": str(overtime.get("type") or ""),
            "source": str(row.get("source") or ""),
            "note": str(row.get("manual_note") or overtime.get("note") or ""),
        })
    out.sort(key=lambda item: (item["date"], 0 if item["slot"] == "RANO" else 1))
    return out


def _employee_inconsistencies(
    login: str,
    year: int,
    month: int,
    *,
    decisions: list[dict] | None = None,
    records: list[dict] | None = None,
) -> list[dict]:
    if decisions is None:
        try:
            import profile_attendance_finalize_runtime as final
            decisions = [
                dict(case)
                for case in final._all_decisions(year, month)
                if str(case.get("login") or "").strip().casefold() == str(login or "").strip().casefold()
            ]
        except Exception:
            decisions = []
    if records is None:
        records = [dict(row) for row in attendance_service.month_records(login, year, month)]

    issues: list[dict] = []
    for case in decisions:
        day_text = str(case.get("date") or "")[:10]
        slot = str(case.get("slot") or "")
        if case.get("is_conflict"):
            text = str(case.get("decision_label") or "Sprzeczne dane obecności")
            kind = "Konflikt"
        else:
            text = str(case.get("decision_label") or "Pozycja wymaga decyzji")
            kind = "Do decyzji"
        issues.append({"date": day_text, "slot": slot, "kind": kind, "text": text})

    by_day: dict[str, set[str]] = {}
    for row in records:
        if row.get("synthetic"):
            continue
        day_text = str(row.get("date") or "")[:10]
        slot = str(row.get("slot") or "").strip().upper()
        if day_text and slot in {attendance_service.RANO, attendance_service.POPO}:
            by_day.setdefault(day_text, set()).add(slot)
    for day_text, slots in sorted(by_day.items()):
        if {attendance_service.RANO, attendance_service.POPO}.issubset(slots):
            issues.append({
                "date": day_text,
                "slot": "RANO + POPO",
                "kind": "Duplikat zmiany",
                "text": "Ten sam pracownik ma wpis na obu zmianach tego dnia.",
            })

    unique: list[dict] = []
    seen: set[tuple[str, str, str, str]] = set()
    for issue in issues:
        key = (
            str(issue.get("date") or ""), str(issue.get("slot") or ""),
            str(issue.get("kind") or ""), str(issue.get("text") or ""),
        )
        if key not in seen:
            seen.add(key)
            unique.append(issue)
    unique.sort(key=lambda item: (str(item.get("date") or ""), str(item.get("kind") or "")), reverse=True)
    return unique


def _export_attendance_xlsx(path: str | Path, login: str, year: int, month: int) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    output = Path(path)
    rows = _attendance_export_rows(login, year, month)
    user = workforce_profile_service.get_user(login) or {"login": login}
    display_name = workforce_profile_service.display_name(user)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ewidencja"
    ws.append(["Pracownik", display_name])
    ws.append(["Login", login])
    ws.append(["Miesiąc", f"{year:04d}-{month:02d}"])
    ws.append([])
    headers = [
        "Data", "Zmiana", "Pierwsze logowanie", "Status", "Dniówka",
        "Nieobecność", "Nadgodziny [h]", "Typ nadgodzin", "Źródło", "Uwagi",
    ]
    ws.append(headers)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([
            row["date"], row["slot"], row["first_login"], row["status"], row["day_value"],
            row["absence"], row["overtime_hours"], row["overtime_type"], row["source"], row["note"],
        ])
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:J{ws.max_row}"
    widths = (12, 11, 18, 18, 10, 15, 16, 18, 18, 36)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    summary = attendance_service.summary_for_month(login, year, month)
    absence_days: dict[str, set[str]] = {code: set() for code in ("L4", "ŚW", "NN", "UR", "UŻ")}
    for row in rows:
        for code in [part.strip() for part in str(row["absence"]).split(",") if part.strip()]:
            if code in absence_days:
                absence_days[code].add(row["date"])

    sm = wb.create_sheet("Podsumowanie")
    sm.append(["Pracownik", display_name])
    sm.append(["Miesiąc", f"{year:04d}-{month:02d}"])
    sm.append([])
    sm.append(["Pozycja", "Wartość"])
    for cell in sm[4]:
        cell.font = Font(bold=True)
    summary_rows = [
        ("Dniówki", float(summary.get("days") or 0.0)),
        ("Soboty", float(summary.get("saturday_days") or 0.0)),
        ("Nadgodziny [h]", float(summary.get("overtime_hours") or 0.0)),
        ("L4 [dni]", len(absence_days["L4"])),
        ("ŚW [dni]", len(absence_days["ŚW"])),
        ("NN [dni]", len(absence_days["NN"])),
        ("UR [dni]", len(absence_days["UR"])),
        ("UŻ [dni]", len(absence_days["UŻ"])),
        ("Braki", float(summary.get("missing") or 0.0)),
        ("Do decyzji", float(summary.get("pending") or 0.0)),
    ]
    for item in summary_rows:
        sm.append(list(item))
    sm.column_dimensions["A"].width = 24
    sm.column_dimensions["B"].width = 16

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def _decorate_employee_attendance(win, nb: ttk.Notebook, login: str) -> None:
    attendance_id = _tabs(nb).get("Obecność")
    if not attendance_id:
        return
    try:
        att = win.nametowidget(attendance_id)
    except Exception:
        return

    user = workforce_profile_service.get_user(login) or {}
    user_id = str(user.get("user_id") or "")
    month_var = tk.StringVar(value=date.today().strftime("%Y-%m"))

    tools = ttk.LabelFrame(att, text="Miesięczne zestawienie", padding=10)
    tools.grid(row=6, column=0, columnspan=3, sticky="ew", pady=(18, 8))
    tools.columnconfigure(1, weight=1)
    ttk.Label(tools, text="Miesiąc:").grid(row=0, column=0, sticky="w")
    month_box = ttk.Combobox(tools, textvariable=month_var, values=_month_choices(), state="readonly", width=10)
    month_box.grid(row=0, column=1, sticky="w", padx=(8, 12))

    def current_login() -> str:
        return _resolve_login(user_id, login)

    def export_month() -> None:
        try:
            year, month = _month_tuple(month_var.get())
        except Exception as exc:
            messagebox.showerror("Eksport obecności", str(exc), parent=win)
            return
        employee = current_login()
        safe_login = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in employee) or "pracownik"
        target = filedialog.asksaveasfilename(
            parent=win,
            title="Eksport miesięcznej obecności pracownika",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=f"obecnosc_{safe_login}_{year:04d}-{month:02d}.xlsx",
        )
        if not target:
            return
        try:
            saved = _export_attendance_xlsx(target, employee, year, month)
        except Exception as exc:
            messagebox.showerror("Eksport obecności", f"Nie udało się utworzyć pliku Excel:\n{exc}", parent=win)
            return
        messagebox.showinfo("Eksport obecności", f"Zapisano miesięczne zestawienie:\n{saved}", parent=win)

    ttk.Button(tools, text="Eksportuj miesiąc do Excel", command=export_month).grid(row=0, column=2, sticky="e")
    add_help_button(
        tools,
        "Eksportuje ewidencję tylko tego pracownika i wybranego miesiąca. Plik zawiera dniówki, zmianę, nieobecności, nadgodziny i podsumowanie.",
        row=0, column=3, padx=(6, 0), sticky="w",
    )

    issues_box = ttk.LabelFrame(att, text="⚠ Niezgodności", padding=10)
    issues_box.grid(row=7, column=0, columnspan=3, sticky="ew", pady=(4, 8))
    issue_top = ttk.Frame(issues_box)
    issue_top.pack(fill="x", pady=(0, 5))
    ttk.Label(issue_top, text="WM wykrył dane wymagające decyzji lub korekty.").pack(side="left")
    add_help_button(
        issue_top,
        "Sekcja pojawia się tylko, gdy WM wykryje dane wymagające decyzji lub korekty. Po rozwiązaniu ostatniego problemu znika.",
    ).pack(side="left", padx=(6, 0))
    issue_list = ttk.Frame(issues_box)
    issue_list.pack(fill="x")

    def refresh_issues(_event=None) -> None:
        try:
            year, month = _month_tuple(month_var.get())
            issues = _employee_inconsistencies(current_login(), year, month)
        except Exception:
            issues = []
        for child in issue_list.winfo_children():
            child.destroy()
        if not issues:
            issues_box.grid_remove()
            return
        issues_box.grid()
        for issue in issues[:6]:
            text = (
                f"• {issue.get('date') or '—'}  {issue.get('slot') or '—'} — "
                f"{issue.get('kind')}: {issue.get('text')}"
            )
            ttk.Label(issue_list, text=text, wraplength=850, justify="left").pack(anchor="w", pady=1)
        if len(issues) > 6:
            ttk.Label(issue_list, text=f"… oraz {len(issues) - 6} kolejnych pozycji.").pack(anchor="w", pady=(3, 0))

    month_box.bind("<<ComboboxSelected>>", refresh_issues, add="+")
    refresh_issues()

    try:
        root = win.winfo_toplevel()
        bind_id = root.bind("<<ProfileDataUpdated>>", lambda _e: win.after_idle(refresh_issues), add="+")
    except Exception:
        root, bind_id = None, None
    if root is not None and bind_id:
        def cleanup(event) -> None:
            if event.widget is win:
                try:
                    root.unbind("<<ProfileDataUpdated>>", bind_id)
                except Exception:
                    pass
        win.bind("<Destroy>", cleanup, add="+")


def _simplify_tabs(win, nb: ttk.Notebook) -> None:
    tabs = _tabs(nb)
    basic_id = tabs.get("Dane")
    perms_id = tabs.get("Uprawnienia")
    history_id = tabs.get("Historia")
    more_id = tabs.get("Więcej")
    if not more_id:
        return
    try:
        more = win.nametowidget(more_id)
        if basic_id:
            nb.tab(basic_id, text="Podstawowe")
        if perms_id:
            nb.hide(perms_id)
        if history_id:
            nb.hide(history_id)
    except Exception:
        return

    access = ttk.LabelFrame(more, text="Historia i administracja", padding=12)
    access.pack(fill="x", pady=(14, 0))
    ttk.Label(access, text="Rzadziej używane informacje są schowane, ale nadal dostępne z tego miejsca.").pack(anchor="w", pady=(0, 10))
    actions = ttk.Frame(access)
    actions.pack(fill="x")

    def open_history() -> None:
        if not history_id:
            return
        try:
            nb.add(history_id, text="Historia")
            nb.select(history_id)
        except Exception:
            pass

    def open_permissions() -> None:
        if not perms_id:
            return
        try:
            nb.add(perms_id, text="Uprawnienia")
            nb.select(perms_id)
        except Exception:
            pass

    if history_id:
        ttk.Button(actions, text="Historia pracownika", command=open_history).pack(side="left")
        add_help_button(
            actions,
            "Otwiera pełną historię zmian pracownika w czytelnej formie. Historia pozostaje tylko do odczytu i zachowuje wcześniejsze wpisy audytowe.",
        ).pack(side="left", padx=(6, 16))
    if perms_id:
        ttk.Button(actions, text="Uprawnienia pracownika", command=open_permissions).pack(side="left")
        add_help_button(
            actions,
            "Otwiera indywidualne wyłączenia modułów dla tego pracownika. Po przejściu do innej zakładki Uprawnienia ponownie zostaną schowane pod Więcej.",
        ).pack(side="left", padx=(6, 0))

    def tab_changed(_event=None) -> None:
        try:
            selected = nb.select()
            for tab_id in (history_id, perms_id):
                if tab_id and selected != tab_id and tab_id in nb.tabs():
                    nb.hide(tab_id)
        except Exception:
            pass

    nb.bind("<<NotebookTabChanged>>", tab_changed, add="+")


def _polish_editor(win, nb: ttk.Notebook) -> None:
    try:
        win.minsize(900, 620)
    except Exception:
        pass
    try:
        for tab_id in nb.tabs():
            frame = win.nametowidget(tab_id)
            frame.columnconfigure(1, weight=1)
    except Exception:
        pass


def _refresh_open_profile_views(owner) -> None:
    try:
        root = owner.winfo_toplevel()
    except Exception:
        return
    try:
        root.event_generate("<<ProfileDataUpdated>>", when="tail")
    except Exception:
        pass
    for widget in _walk(root):
        callback = None
        if widget.__class__.__name__ == "ForemanProfilePanel":
            callback = getattr(widget, "refresh_data", None)
        elif widget.__class__.__name__ == "ProfileView":
            callback = getattr(widget, "_refresh_view", None)
        if callable(callback):
            try:
                widget.after_idle(callback)
            except Exception:
                pass


def _postprocess(win, login: str) -> None:
    nb = _notebook(win)
    if nb is None:
        return
    _decorate_leave_year(win, nb, login)
    _decorate_history(win, nb, login)
    _decorate_employee_attendance(win, nb, login)
    _simplify_tabs(win, nb)
    _polish_editor(win, nb)


def install() -> None:
    global _INSTALLED
    if _INSTALLED:
        return
    import profile_foreman_edit_runtime as edit_runtime

    if getattr(edit_runtime, "_wm_employee_editor_finish_v1", False):
        _INSTALLED = True
        return

    original = edit_runtime.open_employee_editor

    def open_employee_editor(owner, login: str, *, initial_tab: str = "Dane", on_saved=None) -> None:
        key = _employee_key(owner, login)
        existing = _OPEN_EMPLOYEE_WINDOWS.get(key)
        if existing is not None:
            if _raise_window(existing, initial_tab):
                return
            _OPEN_EMPLOYEE_WINDOWS.pop(key, None)

        try:
            root = owner.winfo_toplevel()
        except Exception:
            root = owner
        before = _all_toplevels(root)

        def saved() -> None:
            try:
                if callable(on_saved):
                    on_saved()
            finally:
                _refresh_open_profile_views(owner)

        original(owner, login, initial_tab=initial_tab, on_saved=saved)
        created = list(_all_toplevels(root) - before)
        if not created:
            return
        win = created[-1]
        _OPEN_EMPLOYEE_WINDOWS[key] = win
        _postprocess(win, login)
        _select_tab(win, initial_tab)

        def cleanup(event) -> None:
            if event.widget is win and _OPEN_EMPLOYEE_WINDOWS.get(key) is win:
                _OPEN_EMPLOYEE_WINDOWS.pop(key, None)

        win.bind("<Destroy>", cleanup, add="+")

    edit_runtime.open_employee_editor = open_employee_editor
    edit_runtime._wm_employee_editor_finish_v1 = True
    _INSTALLED = True


__all__ = ["install", "_attendance_export_rows", "_employee_inconsistencies", "_export_attendance_xlsx"]
