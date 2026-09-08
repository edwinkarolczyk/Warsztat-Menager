# version: 1.0
"""Optymalizacja miesięcznego eksportu obecności i przygotowanie arkusza do druku."""
from __future__ import annotations

from calendar import monthrange
from datetime import date
from pathlib import Path
from typing import Any

from services import attendance_service, leave_workflow_service, workforce_profile_service

_INSTALLED = False


def _matches_export_user(row: dict, login: str) -> bool:
    """Dopasuj nieobecność po user_id z fallbackiem do bieżącego/starego loginu."""
    user = workforce_profile_service.get_user(login) or {}
    uid_key = str(user.get("user_id") or "").strip().casefold()
    login_keys = {
        str(login or "").strip().casefold(),
        str(user.get("login") or "").strip().casefold(),
    }
    login_keys.discard("")

    row_uid = str(row.get("user_id") or "").strip().casefold()
    if uid_key and row_uid == uid_key:
        return True
    row_logins = {
        str(row.get("login") or "").strip().casefold(),
        str(row.get("login_snapshot") or "").strip().casefold(),
    }
    row_logins.discard("")
    return bool(login_keys.intersection(row_logins))


def _month_absence_index(login: str, year: int, month: int) -> dict[str, list[dict]]:
    """Wczytaj aktywne nieobecności raz i pogrupuj tylko wybrany miesiąc."""
    prefix = f"{year:04d}-{month:02d}-"
    try:
        leaves = leave_workflow_service.read_leaves()
    except Exception:
        leaves = []

    by_day: dict[str, list[dict]] = {}
    for raw in leaves:
        if not isinstance(raw, dict) or not _matches_export_user(raw, login):
            continue
        day_text = str(raw.get("date") or "")[:10]
        if not day_text.startswith(prefix):
            continue
        by_day.setdefault(day_text, []).append(dict(raw))
    return by_day


def _absence_code(value: Any) -> str:
    raw = str(value or "").strip().upper().replace("-", "_").replace(" ", "_")
    mapping = {
        "SW": "ŚW", "ŚW": "ŚW", "SILA_WYZSZA": "ŚW", "SIŁA_WYŻSZA": "ŚW",
        "SILA_WYZSZA_50": "ŚW", "URLOP": "UR", "UR": "UR",
        "URLOP_WYPOCZYNKOWY": "UR", "UZ": "UŻ", "UŻ": "UŻ",
        "URLOP_NA_ZADANIE": "UŻ", "L4": "L4", "NN": "NN",
    }
    return mapping.get(raw, raw)


def _absence_codes(row: dict | None, leaves: list[dict]) -> list[str]:
    out: list[str] = []
    reason = _absence_code((row or {}).get("reason"))
    if reason:
        out.append(reason)
    for leave in leaves:
        code = _absence_code(leave.get("type"))
        if code and code not in out:
            out.append(code)
    return out


def _attendance_status_text(row: dict) -> str:
    labels = {
        attendance_service.STATUS_PRESENT: "Obecny",
        attendance_service.STATUS_PENDING_LATE: "Późne logowanie",
        attendance_service.STATUS_MISSING: "Brak",
        attendance_service.STATUS_EXCUSED: "Nieobecność",
        attendance_service.STATUS_SATURDAY: "Sobota",
        attendance_service.STATUS_PLANNED: "Plan",
    }
    status = str(row.get("status") or "")
    return labels.get(status, status or "—")


def _attendance_export_rows(login: str, year: int, month: int) -> list[dict]:
    """Przygotuj miesiąc bez N+1 odczytów leaves.json."""
    records = [dict(row) for row in attendance_service.month_records(login, year, month)]
    absences_by_day = _month_absence_index(login, year, month)

    existing_days = {str(row.get("date") or "")[:10] for row in records}
    for day_no in range(1, monthrange(year, month)[1] + 1):
        day_text = date(year, month, day_no).isoformat()
        if day_text not in existing_days and _absence_codes(None, absences_by_day.get(day_text, [])):
            records.append({
                "date": day_text,
                "slot": "",
                "status": attendance_service.STATUS_EXCUSED,
                "day_value": 0.0,
                "reason": "",
                "synthetic": True,
            })

    out: list[dict] = []
    for row in records:
        day_text = str(row.get("date") or "")[:10]
        overtime = row.get("overtime") if isinstance(row.get("overtime"), dict) else {}
        first_login = str(row.get("first_login_ts") or row.get("logged_ts") or "")
        if "T" in first_login:
            first_login = first_login.split("T", 1)[1][:8]
        codes = _absence_codes(row, absences_by_day.get(day_text, []))
        out.append({
            "date": day_text,
            "slot": str(row.get("slot") or "—"),
            "first_login": first_login or "—",
            "status": _attendance_status_text(row),
            "day_value": float(row.get("day_value") or 0.0),
            "absence": ", ".join(codes),
            "overtime_hours": float(overtime.get("hours") or 0.0),
            "overtime_type": str(overtime.get("type") or ""),
            "source": str(row.get("source") or ""),
            "note": str(row.get("manual_note") or overtime.get("note") or ""),
        })
    out.sort(key=lambda item: (item["date"], 0 if item["slot"] == "RANO" else 1))
    return out


def _prepare_xlsx_for_print(path: str | Path) -> Path:
    """Ustaw zapisany miesięczny Excel tak, aby dało się go od razu wydrukować."""
    output = Path(path)
    try:
        from openpyxl import load_workbook

        wb = load_workbook(output)
        ws = wb["Ewidencja"]
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "5:5"
        ws.print_area = f"A1:J{ws.max_row}"
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5

        if "Podsumowanie" in wb.sheetnames:
            sm = wb["Podsumowanie"]
            sm.page_setup.orientation = "portrait"
            sm.page_setup.fitToWidth = 1
            sm.page_setup.fitToHeight = 0
            sm.sheet_properties.pageSetUpPr.fitToPage = True
            sm.print_area = f"A1:B{sm.max_row}"
            sm.print_options.horizontalCentered = True
        wb.save(output)
    except Exception as exc:
        print(f"[WM-DBG][PROFILE][WARN] print-ready xlsx setup failed: {exc!r}")
    return output


def install() -> None:
    global _INSTALLED
    import profile_employee_editor_finish_runtime as target

    if getattr(target, "_wm_attendance_export_opt_v1", False):
        _INSTALLED = True
        return

    original_export_xlsx = target._export_attendance_xlsx

    def export_xlsx_print_ready(path: str | Path, login: str, year: int, month: int) -> Path:
        output = original_export_xlsx(path, login, year, month)
        return _prepare_xlsx_for_print(output)

    target._attendance_export_rows = _attendance_export_rows
    target._export_attendance_xlsx = export_xlsx_print_ready
    target._wm_attendance_export_opt_v1 = True
    _INSTALLED = True


__all__ = [
    "_attendance_export_rows",
    "_month_absence_index",
    "_prepare_xlsx_for_print",
    "install",
]
